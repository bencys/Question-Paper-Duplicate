"""Build a new question paper as a Word document."""

from __future__ import annotations

import base64
from io import BytesIO

from docx import Document
from docx.shared import Cm, Pt

from images import to_png


def _picture(data_uri: str) -> BytesIO | None:
    """The picture as PNG bytes, or None when it cannot be read; a paper is still built without it."""
    try:
        return BytesIO(to_png(base64.b64decode(data_uri.split(",", 1)[1])))
    except Exception:
        return None


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    cols = max((len(r) for r in rows), default=0)
    if not rows or not cols:
        return
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            table.cell(r, c).text = text


def build_paper(title: str, questions: list[dict]) -> bytes:
    """Numbered questions with their pictures, tables and MCQ options, as a .docx file."""
    doc = Document()
    doc.add_heading(title, level=1)
    for i, q in enumerate(questions, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = bool(q.get("options"))
        p.add_run(f"{i}. ").bold = True
        p.add_run(q["text"])
        for src in q.get("images") or []:
            picture = _picture(src)
            if picture:
                doc.add_picture(picture, width=Cm(9))
        for table in q.get("tables") or []:
            _add_table(doc, table)
        option_images = q.get("option_images") or []
        for k, option in enumerate(q.get("options") or []):
            o = doc.add_paragraph()
            o.paragraph_format.left_indent = Cm(1)
            o.paragraph_format.space_after = Pt(0)
            o.add_run(option)
            for src in option_images[k] if k < len(option_images) else []:
                picture = _picture(src)
                if picture:
                    o.add_run().add_picture(picture, height=Cm(2.5))
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
