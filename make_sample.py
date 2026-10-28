"""Create a sample question paper for a quick demo."""

from docx import Document

doc = Document()
doc.add_heading("Sample Internal Exam — Computer Science", level=1)
doc.add_paragraph("Time: 3 Hours    Maximum Marks: 70")
doc.add_paragraph("Answer any five questions. All questions carry equal marks.")

qs = [
    "1. Explain the difference between RAM and ROM.",
    "2. What is an operating system? Describe its main functions.",
    "3. Explain the difference between RAM and ROM.",
    "4. Diference between RAM and ROM explain.",
    "5. Describe how RAM differs from ROM.",
    "6. Define a compiler and an interpreter.",
    "7. Write a short note on TCP/IP model.",
    "8. What is the TCP IP model? Write a short note.",
    "9. List any four input devices.",
    "10. Mention four input devices used with a computer.",
    "11. What is a database management system?",
    "12. Explain polymorphism in object oriented programming.",
    "13. Explain polymorfism in object oriented programing.",
    "14. Convert the binary number 1011 to decimal.",
]
for q in qs:
    doc.add_paragraph(q)

path = "sample_question_paper.docx"
doc.save(path)
print("Wrote", path)
